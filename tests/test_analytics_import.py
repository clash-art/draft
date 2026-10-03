import io,sys,unittest,base64
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from analytics_import import read_export
class ExportTests(unittest.TestCase):
 def test_xlsx_numbers_dates_preamble(self):
  from openpyxl import Workbook
  from datetime import datetime
  w=Workbook();s=w.active;s.append(['公众号数据']);s.append(['日期','文章标题','阅读次数']);s.append([datetime(2026,9,19),'文章',120]);b=io.BytesIO();w.save(b)
  rows,_=read_export(b.getvalue(),'数据.xlsx')
  self.assertEqual(rows,[{'ref_date':'2026-09-19','title':'文章','int_page_read_count':120}])
 def test_html_xls(self):
  raw='<table><tr><th>标题</th><th>阅读次数</th></tr><tr><td>文章</td><td>25</td></tr></table>'.encode()
  self.assertEqual(read_export(raw,'数据.xls')[0][0]['int_page_read_count'],'25')
 def test_csv_and_invalid(self):
  self.assertEqual(read_export('标题,阅读次数\n文章,3'.encode(),'数据.csv')[0][0]['title'],'文章')
  with self.assertRaises(ValueError):read_export(b'junk','bad.xls')
 def test_multiple_sheets_not_silently_summed(self):
  from openpyxl import Workbook
  w=Workbook();w.active.append(['标题','阅读次数']);w.active.append(['一',2]);s=w.create_sheet();s.append(['标题','阅读次数']);s.append(['二',2]);b=io.BytesIO();w.save(b)
  with self.assertRaisesRegex(ValueError,'多个'):read_export(b.getvalue(),'数据.xlsx')
 def test_real_binary_xls(self):
  raw=(Path(__file__).parent/'fixtures/analytics.xls').read_bytes()
  rows,_=read_export(raw,'微信数据.xls')
  self.assertEqual(rows[0]['int_page_read_count'],123)
  self.assertEqual(rows[0]['title'],'XLS 测试')

 def test_workspace_excel_snapshot_and_incremental(self):
  import tempfile
  from workspace import Workspace
  raw=(Path(__file__).parent/'fixtures/analytics.xls').read_bytes()
  with tempfile.TemporaryDirectory() as directory:
   w=Workspace(Path(directory),lambda:None)
   data={'name':'data.xls','data':base64.b64encode(raw).decode()}
   report=w.dispatch('/api/analytics/import',data)
   self.assertEqual(report['totals'],{})
   self.assertEqual(report['source'],'import')
   report=w.dispatch('/api/analytics/import',dict(data,mode='incremental'))
   self.assertEqual(report['totals']['int_page_read_count'],123)
 def test_wechat_multisection_report_keeps_uv_and_channels(self):
  from analytics_import import read_overview_export
  raw=(Path(__file__).parent/'fixtures/article-report.xls').read_bytes()
  report=read_overview_export(raw,'report.xls')
  self.assertEqual(report['overview'][0],['阅读(人)',117])
  self.assertEqual(report['records'][0]['int_page_read_user'],57)
  self.assertNotIn('int_page_read_count',report['records'][0])
  self.assertEqual(report['records'][1]['channel'],'Recommendation')
  self.assertEqual(report['totals'],{})
  self.assertEqual(len(report['sections']),3)
