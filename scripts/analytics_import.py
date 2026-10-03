"""Read Excel/CSV exports without changing the original file."""
import csv, io, zipfile
from datetime import date, datetime
from pathlib import Path
from wechat import ALIASES

MAX_BYTES=4*1024*1024

def normalized(value):
    if value is None:return ''
    if isinstance(value,(date,datetime)):return value.isoformat()[:10]
    if isinstance(value,float) and value.is_integer():return int(value)
    return value.strip() if isinstance(value,str) else value

def records(matrix):
    # Exports may start with a report title or date-range preamble.
    header=None;result=[]
    for index,row in enumerate(matrix):
        if index>50000:raise ValueError('数据表不能超过 50000 行')
        values=[normalized(v) for v in row]
        if not header:
            keys=[ALIASES.get(str(v).strip(),str(v).strip()) for v in values]
            if any(k in ('int_page_read_count','share_count','add_to_fav_count','int_page_read_user') for k in keys):
                header=keys
                nonempty=[k for k in keys if k]
                if len(nonempty)!=len(set(nonempty)):raise ValueError('数据表含重复列名，请整理后导入')
            elif index>=30:return []
            continue
        if not any(v!='' for v in values):continue
        row={k:v for k,v in zip(header,values) if k}
        if row.get('ref_date') in ('合计','总计') or row.get('title') in ('合计','总计'):continue
        result.append(row)
    return result

def read_export(raw,name):
    if not raw or len(raw)>MAX_BYTES:raise ValueError('文件为空或超过 4 MB')
    suffix=Path(name).suffix.lower()
    if suffix not in ('.xls','.xlsx','.csv'):raise ValueError('请选择 XLS、XLSX 或 CSV 文件')
    tables=[]
    try:
        if raw.startswith(bytes.fromhex('D0CF11E0A1B11AE1')):
            import xlrd
            book=xlrd.open_workbook(file_contents=raw,on_demand=True)
            try:
                for sheet in book.sheets():
                    if sheet.nrows>50001 or sheet.ncols>256:raise ValueError('数据表过大')
                    matrix=[]
                    for r in range(sheet.nrows):
                        matrix.append([xlrd.xldate_as_datetime(c.value,book.datemode) if c.ctype==xlrd.XL_CELL_DATE else c.value for c in sheet.row(r)])
                    tables.append((sheet.name,records(matrix)))
            finally:book.release_resources()
        elif raw.startswith(b'PK'):
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(i.file_size for i in archive.infolist())>40*1024*1024:raise ValueError('Excel 解压后过大')
            book=load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
            try:
                for sheet in book:
                    if (sheet.max_column or 0)>256 or (sheet.max_row or 0)>50001:raise ValueError('数据表过大')
                    tables.append((sheet.title,records(sheet.iter_rows(values_only=True))))
            finally:book.close()
        else:
            text=None
            for encoding in (('utf-16',) if raw.startswith((b'\xff\xfe',b'\xfe\xff')) else ('utf-8-sig','gb18030')):
                try:text=raw.decode(encoding);break
                except UnicodeError:pass
            if text is None:raise ValueError('无法识别文件编码')
            if '<table' in text.lower():
                from bs4 import BeautifulSoup
                for n,table in enumerate(BeautifulSoup(text,'html.parser').find_all('table')):
                    tables.append((f'表格 {n+1}',records([[c.get_text(strip=True) for c in row.find_all(['td','th'])] for row in table.find_all('tr')])))
            else:
                delimiter='\t' if '\t' in text.splitlines()[0] else ','
                tables=[('数据',records(csv.reader(io.StringIO(text),delimiter=delimiter)))]
    except ValueError:raise
    except Exception:raise ValueError('无法解析表格，请重新导出未加密的 XLS、XLSX 或 CSV 文件') from None
    usable=[(name,rows) for name,rows in tables if rows]
    if not usable:raise ValueError('未识别到阅读次数、阅读人数、分享次数或收藏次数等数据列')
    if len(usable)>1:raise ValueError('文件包含多个数据工作表，请保留要分析的一张表后导入，避免重复统计')
    return usable[0][1],usable[0][0]

def read_overview_export(raw,name):
    """Recognize the multi-section WeChat article report, preserving reported UV."""
    if len(raw)>MAX_BYTES:raise ValueError('文件超过 4 MB')
    if not raw.startswith(bytes.fromhex('D0CF11E0A1B11AE1')):return None
    import xlrd
    book=xlrd.open_workbook(file_contents=raw)
    if book.nsheets!=1:return None
    sheet=book.sheet_by_index(0)
    if sheet.nrows>50001 or sheet.ncols>256:raise ValueError('数据表过大')
    matrix=[[normalized(v) for v in sheet.row_values(i)] for i in range(sheet.nrows)]
    if not any('Data overview' in r or '数据概览' in r for r in matrix):return None
    title=next((str(v) for r in matrix[:2] for v in r if v!=''),'')
    labels={'Data overview':'数据概览','View conversion':'阅读转化','阅读数据趋势明细':'每日渠道明细','Statistics by gender':'性别分布','Statistics by age':'年龄分布','Statistics by location':'地域分布'}
    sections=[];current=None
    for row in matrix[1:]:
        row=list(row)
        while row and row[0]=='':row.pop(0)
        while row and row[-1]=='':row.pop()
        if not row:continue
        if len(row)==1 and row[0] in labels:
            current={'name':labels[row[0]],'columns':[],'rows':[]};sections.append(current);continue
        if current:
            if not current['columns']:current['columns']=row
            else:current['rows'].append(row)
    overview=next((s['rows'] for s in sections if s['name']=='数据概览'),[])
    trend=next((s['rows'] for s in sections if s['name']=='每日渠道明细'),[])
    rows=[{'title':title,'ref_date':r[0],'channel':r[1],'int_page_read_user':r[2],'share_user':r[3]} for r in trend if len(r)>=4]
    if not overview or not rows:return None
    return {'rows':len(rows),'records':rows,'totals':{},'overview':overview,'sections':sections,'report_title':title,'format':'wechat_article_report','complete':True,'source':'import','notes':['概览采用微信报表原值；阅读人数和分享人数不跨日期或渠道相加。','All 为当日整体，不能与渠道明细重复相加。']}
